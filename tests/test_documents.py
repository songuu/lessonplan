import io
import zipfile

import pytest

from app.integrations.documents import DocumentAdapter, DocumentError
from app.integrations.storage import LocalStorage, StorageError

CONTENT = "# 分数教案\n\n中文与 IPA /ˈtiːtʃə/。\\(\\frac{1}{2}+x^2\\)\n\n- 观察\n- 解释\n\n| 环节 | 时长 |\n| --- | --- |\n| 操作 | 10分钟 |"


def test_docx_readback_preserves_table_ipa_and_native_math():
    from docx import Document
    adapter = DocumentAdapter()
    data = adapter.export(CONTENT, "分数", "docx")
    doc = Document(io.BytesIO(data))
    assert "分数教案" in "\n".join(p.text for p in doc.paragraphs)
    assert "/ˈtiːtʃə/" in "\n".join(p.text for p in doc.paragraphs)
    assert doc.tables[0].cell(1, 0).text == "操作"
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        xml = archive.read("word/document.xml").decode()
    assert "m:oMath" in xml
    parsed = adapter.parse_upload("分数.docx", data)
    assert "中文" in parsed["text"] and "操作" in parsed["text"]
    assert parsed["style_summary"]["heading_count"] >= 1


def test_pdf_readback_has_chinese_and_formula():
    from pypdf import PdfReader
    data = DocumentAdapter().export(CONTENT, "分数", "pdf")
    reader = PdfReader(io.BytesIO(data))
    assert len(reader.pages) >= 1
    text = "\n".join(page.extract_text() for page in reader.pages)
    assert "分数" in text and "中文" in text
    assert "x" in text
    parsed = DocumentAdapter().parse_upload("lesson.pdf", data)
    assert "中文" in parsed["text"]


def test_upload_validates_signature_size_and_zip_expansion():
    adapter = DocumentAdapter(max_upload_bytes=1024, max_expanded_bytes=2048)
    with pytest.raises(DocumentError):
        adapter.parse_upload("bad.pdf", b"not a pdf")
    with pytest.raises(DocumentError):
        adapter.parse_upload("bad.docx", b"x" * 2048)
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "a" * 3000)
    with pytest.raises(DocumentError) as error:
        adapter.parse_upload("bomb.docx", output.getvalue())
    assert error.value.code == "upload_expansion_limit"


def test_local_storage_keys_never_accept_paths(tmp_path):
    storage = LocalStorage(tmp_path)
    key = storage.put(b"hello", suffix=".docx")
    assert storage.get(key) == b"hello"
    with pytest.raises(StorageError):
        storage.get("../sensitive.txt")
    with pytest.raises(StorageError):
        storage.put(b"hello", suffix="../../a")
    storage.delete(key)
    with pytest.raises(StorageError):
        storage.get(key)


def test_omml_arguments_are_valid_math_children_and_complex_source_is_preserved():
    from lxml import etree
    adapter = DocumentAdapter()
    data = adapter.export(r"\(\frac{1}{2}+\sqrt{x_1^2}\) and \(\unsupported{x}\)", "公式", "docx")
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        tree = etree.fromstring(archive.read("word/document.xml"))
    namespace = {"m": "http://schemas.openxmlformats.org/officeDocument/2006/math"}
    assert tree.xpath("count(//m:f)", namespaces=namespace) == 1
    assert tree.xpath("count(//m:num/m:e)", namespaces=namespace) == 0
    assert tree.xpath("count(//m:sSubSup)", namespaces=namespace) == 1
    assert "\\unsupported{x}" in "".join(tree.itertext())
    assert "math_source_preserved" in adapter.warnings


def test_cos_adapter_uses_opaque_keys_and_never_exposes_sdk_details():
    from types import SimpleNamespace

    from app.integrations.storage import COSStorage

    class Client:
        def put_object(self, **kwargs):
            self.kwargs = kwargs

        def get_object(self, **kwargs):
            raise ValueError("sensitive-token-in-sdk-error")

    client = Client()
    storage = COSStorage(SimpleNamespace(cos_bucket="example-bucket"), client=client)
    key = storage.put(b"data", suffix=".pdf", content_type="application/pdf")
    assert client.kwargs["Key"] == key
    assert client.kwargs["ContentType"] == "application/pdf"
    with pytest.raises(StorageError) as error:
        storage.get(key)
    assert "sensitive" not in str(error.value)


def test_export_rejects_dangerous_tex_and_upload_rejects_legacy_doc():
    adapter = DocumentAdapter()
    with pytest.raises(DocumentError):
        adapter.export(r"\(\write18{execute}\)", "unsafe", "pdf")
    with pytest.raises(DocumentError) as error:
        adapter.parse_upload("old.doc", b"old-word-file")
    assert error.value.code == "unsupported_upload_format"
