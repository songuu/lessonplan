"""Bounded upload parsing and independently readable DOCX/PDF exports."""
from __future__ import annotations

import io
import os
import re
import threading
import zipfile
from functools import lru_cache
from html import escape
from pathlib import Path

from app.ai.quality import ModuleQualityGate, SafePatchProcessor

from .formula import FormulaUnsupported, formula_segments, math_box, omml, parse_formula


class DocumentError(Exception):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code, self.message = code, message


_FONT_LOCK = threading.Lock()


@lru_cache(maxsize=1)
def _pdf_fonts():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfbase.ttfonts import TTFont
    cjk_candidates = [os.getenv("LESSON_CJK_FONT", ""), "C:/Windows/Fonts/simsun.ttc",
                      "C:/Windows/Fonts/msyh.ttc", "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc",
                      "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
                      "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttf"]
    latin_candidates = [os.getenv("LESSON_LATIN_FONT", ""), "C:/Windows/Fonts/arial.ttf",
                        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
                        "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"]
    warnings = []
    with _FONT_LOCK:
        cjk, latin = "STSong-Light", "Helvetica"
        for name, candidates in (("LessonCJK", cjk_candidates), ("LessonLatin", latin_candidates)):
            for candidate in candidates:
                if not candidate or not Path(candidate).is_file():
                    continue
                try:
                    pdfmetrics.registerFont(TTFont(name, candidate, subfontIndex=0))
                    if name == "LessonCJK":
                        cjk = name
                    else:
                        latin = name
                    break
                except Exception:
                    continue
        if cjk == "STSong-Light":
            pdfmetrics.registerFont(UnicodeCIDFont(cjk))
            warnings.append("pdf_cjk_cid_fallback")
        if latin == "Helvetica":
            warnings.append("pdf_latin_font_fallback")
    return cjk, latin, warnings


def _blocks(content):
    lines = content.splitlines()
    index = 0
    paragraph = []
    while index < len(lines):
        line = lines[index]
        if index + 1 < len(lines) and "|" in line and re.fullmatch(r"\s*\|?\s*:?-{3,}:?\s*(?:\|\s*:?-{3,}:?\s*)+\|?\s*", lines[index + 1]):
            if paragraph:
                yield "paragraph", "\n".join(paragraph)
                paragraph = []
            rows = [[cell.strip() for cell in line.strip().strip("|").split("|")]]
            index += 2
            while index < len(lines) and "|" in lines[index] and lines[index].strip():
                rows.append([cell.strip() for cell in lines[index].strip().strip("|").split("|")])
                index += 1
            yield "table", rows
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        item = re.match(r"^\s*(?:[-*+]\s+|\d+[.)、]\s+)(.+)$", line)
        if heading or item or not line.strip():
            if paragraph:
                yield "paragraph", "\n".join(paragraph)
                paragraph = []
            if heading:
                yield "heading", (len(heading.group(1)), heading.group(2))
            elif item:
                yield "list", item.group(1)
        else:
            paragraph.append(line)
        index += 1
    if paragraph:
        yield "paragraph", "\n".join(paragraph)


def _plain(text):
    text = re.sub(r"\*\*([^*]+)\*\*|__([^_]+)__", lambda m: m.group(1) or m.group(2), text)
    text = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)", r"\1", text)
    text = re.sub(r"\[([^]]+)\]\([^)]*\)", r"\1", text)
    return text.replace("`", "")


class DocumentAdapter:
    def __init__(self, *, max_upload_bytes: int = 20 * 1024 * 1024,
                 max_expanded_bytes: int = 100 * 1024 * 1024, max_pages: int = 200):
        self.max_upload_bytes, self.max_expanded_bytes, self.max_pages = max_upload_bytes, max_expanded_bytes, max_pages
        self.warnings: list[str] = []

    def export(self, content: str, title: str = "教案", format: str = "docx") -> bytes:
        self.warnings = []
        if not isinstance(content, str) or len(content) > 2 * 1024 * 1024:
            raise DocumentError("export_size_limit", "导出正文超出大小限制。")
        # The export adapter is a second safety boundary, including for historical content.
        quality = ModuleQualityGate().check(content, {}, {})
        if not quality["passed"]:
            raise DocumentError("export_unsafe_content", "导出正文包含不安全内容，请先修改。")
        if quality["math_render_state"] == "invalid":
            self.warnings.append("math_source_preserved")
        content, _ = SafePatchProcessor().apply(content, {})
        title = re.sub(r"[\x00-\x1f]", "", str(title))[:200] or "教案"
        if format in {"md", "markdown"}:
            return content.encode("utf-8")
        try:
            if format == "docx":
                return self._docx(content, title)
            if format == "pdf":
                return self._pdf(content, title)
        except DocumentError:
            raise
        except Exception:
            raise DocumentError("document_export_failed", "文档导出失败，请稍后重试。") from None
        raise DocumentError("unsupported_export_format", "仅支持 DOCX、PDF 和 Markdown 导出。")

    def _word_inline(self, paragraph, content):
        from docx.oxml import OxmlElement
        for is_math, value in formula_segments(content):
            if is_math:
                try:
                    node = parse_formula(value)
                    math = OxmlElement("m:oMath")
                    inner = omml(node)
                    # oMath directly contains math elements; m:e is only an argument wrapper.
                    for child in list(inner):
                        math.append(child)
                    paragraph._p.append(math)
                except FormulaUnsupported:
                    paragraph.add_run(r"\(" + value + r"\)")
                    self.warnings.append("math_source_preserved")
            else:
                for fragment in re.split(r"(\*\*[^*]+\*\*|__[^_]+__)", value):
                    if not fragment:
                        continue
                    bold = fragment.startswith("**") or fragment.startswith("__")
                    run = paragraph.add_run(_plain(fragment[2:-2] if bold else fragment))
                    run.bold = bold

    def _docx(self, content, title):
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Cm, Pt
        document = Document()
        document.core_properties.title = title
        section = document.sections[0]
        section.top_margin = section.bottom_margin = Cm(2)
        section.left_margin = section.right_margin = Cm(2.2)
        for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "List Bullet"):
            style = document.styles[name]
            style.font.name = "Arial"
            style.font.size = Pt(11 if name in {"Normal", "List Bullet"} else 16 if name == "Title" else 14)
            fonts = style.element.get_or_add_rPr().find(qn("w:rFonts"))
            if fonts is None:
                fonts = OxmlElement("w:rFonts")
                style.element.get_or_add_rPr().append(fonts)
            fonts.set(qn("w:eastAsia"), "宋体")
        for kind, value in _blocks(content):
            if kind == "table":
                width = max(len(row) for row in value)
                table = document.add_table(rows=len(value), cols=width)
                table.style = "Table Grid"
                for row_index, row in enumerate(value):
                    for col_index, cell in enumerate(row):
                        self._word_inline(table.cell(row_index, col_index).paragraphs[0], cell)
                # Word repeats the table header when the table spans pages.
                header = OxmlElement("w:tblHeader")
                table.rows[0]._tr.get_or_add_trPr().append(header)
            else:
                if kind == "heading":
                    level, text = value
                    paragraph = document.add_paragraph(style=f"Heading {min(level, 3)}")
                else:
                    text = value
                    paragraph = document.add_paragraph(style="List Bullet" if kind == "list" else "Normal")
                self._word_inline(paragraph, text)
        document.core_properties.comments = "Export diagnostics: " + ",".join(sorted(set(self.warnings)))
        output = io.BytesIO()
        document.save(output)
        return output.getvalue()

    def _pdf(self, content, title):
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_LEFT
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.platypus import Flowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
        cjk, latin, warnings = _pdf_fonts()
        self.warnings.extend(warnings)
        body = ParagraphStyle("Body", fontName=cjk, fontSize=10.5, leading=17, spaceAfter=7,
                              wordWrap="CJK", alignment=TA_LEFT)
        heading_styles = {i: ParagraphStyle(f"H{i}", parent=body, fontSize=17 - i, leading=24 - i,
                                            spaceBefore=9, spaceAfter=7, keepWithNext=True) for i in range(1, 7)}

        def rich(text):
            fragments = []
            for value in re.split(r"([\u2e80-\u9fff\uf900-\ufaff\uff00-\uffef]+)", _plain(text)):
                if not value:
                    continue
                font = cjk if re.match(r"[\u2e80-\u9fff\uf900-\ufaff\uff00-\uffef]", value) else latin
                fragments.append(f'<font name="{font}">{escape(value).replace(chr(10), "<br/>")}</font>')
            return "".join(fragments)

        class FormulaFlowable(Flowable):
            def __init__(self, node):
                super().__init__()
                self.box = math_box(node, latin, 12, cjk_font=cjk)
                self.scale = 1

            def wrap(self, available_width, available_height):
                self.scale = min(1, available_width / max(self.box[0], 1))
                self.width = self.box[0] * self.scale
                self.height = (self.box[1] + self.box[2]) * self.scale + 10
                return self.width, self.height

            def draw(self):
                self.canv.saveState()
                self.canv.scale(self.scale, self.scale)
                self.box[3](self.canv, 0, self.box[2] + 4)
                self.canv.restoreState()

        def paragraphs(text, style=body):
            result = []
            for is_math, value in formula_segments(text):
                if is_math:
                    try:
                        result.append(FormulaFlowable(parse_formula(value)))
                    except FormulaUnsupported:
                        result.append(Paragraph(rich(r"\(" + value + r"\)"), style))
                        self.warnings.append("math_source_preserved")
                elif value.strip():
                    result.append(Paragraph(rich(value), style))
            return result or [Spacer(1, 3)]

        story = []
        usable_width = A4[0] - 4 * cm
        for kind, value in _blocks(content):
            if kind == "table":
                width = max(len(row) for row in value)
                rows = [[paragraphs(cell) for cell in row] + [Paragraph("", body)] * (width - len(row)) for row in value]
                table = Table(rows, colWidths=[usable_width / width] * width, repeatRows=1, hAlign="LEFT")
                table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), .5, colors.HexColor("#c3ceda")),
                                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#edf2f7")),
                                           ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                           ("LEFTPADDING", (0, 0), (-1, -1), 7),
                                           ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                                           ("TOPPADDING", (0, 0), (-1, -1), 6),
                                           ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
                story.extend([table, Spacer(1, 9)])
            elif kind == "heading":
                level, text = value
                story.extend(paragraphs(text, heading_styles[level]))
            else:
                story.extend(paragraphs(("• " if kind == "list" else "") + value))
        output = io.BytesIO()

        def page_number(canvas, document):
            canvas.saveState()
            canvas.setFont(latin, 9)
            canvas.setFillColor(colors.HexColor("#64748b"))
            canvas.drawRightString(A4[0] - 2 * cm, 1.2 * cm, str(document.page))
            canvas.restoreState()
        document = SimpleDocTemplate(output, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm,
                                     topMargin=1.7 * cm, bottomMargin=1.8 * cm,
                                     title=title, author="教案助手", subject=",".join(sorted(set(self.warnings))))
        document.build(story, onFirstPage=page_number, onLaterPages=page_number)
        return output.getvalue()

    def parse_upload(self, filename: str, data: bytes) -> dict:
        if not data or len(data) > self.max_upload_bytes:
            raise DocumentError("upload_size_limit", "上传文件为空或超过大小限制。")
        extension = Path(filename).suffix.lower()
        metadata = {"filename": Path(filename).name[:255], "format": extension.lstrip("."), "size_bytes": len(data)}
        try:
            if extension == ".docx":
                if not data.startswith(b"PK\x03\x04"):
                    raise DocumentError("upload_signature_mismatch", "DOCX 文件签名不匹配。")
                self._validate_zip(data)
                from docx import Document
                document = Document(io.BytesIO(data))
                paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
                table_rows = [" | ".join(cell.text for cell in row.cells) for table in document.tables for row in table.rows]
                text = "\n".join(paragraphs + table_rows)
                style = {"heading_count": sum(p.style.name.startswith("Heading") or p.style.name.startswith("标题") for p in document.paragraphs),
                         "table_count": len(document.tables),
                         "list_count": sum(p.style.name.startswith("List") for p in document.paragraphs)}
            elif extension == ".pdf":
                if not data.startswith(b"%PDF-"):
                    raise DocumentError("upload_signature_mismatch", "PDF 文件签名不匹配。")
                from pypdf import PdfReader
                reader = PdfReader(io.BytesIO(data), strict=True)
                if reader.is_encrypted:
                    raise DocumentError("upload_encrypted", "请先移除 PDF 密码后上传。")
                if len(reader.pages) > self.max_pages:
                    raise DocumentError("upload_page_limit", "PDF 页数超过限制。")
                pages = []
                expanded_total = 0
                for page in reader.pages:
                    # Bound decompressed streams before invoking the text interpreter.
                    contents = page.get_contents()
                    if contents is not None:
                        expanded_total += len(contents.get_data())
                        if expanded_total > self.max_expanded_bytes:
                            raise DocumentError("upload_expansion_limit", "PDF 解压后超过安全限制。")
                    pages.append(page.extract_text() or "")
                    if sum(map(len, pages)) > 2 * 1024 * 1024:
                        raise DocumentError("upload_text_limit", "文档正文超过安全限制。")
                text = "\n".join(pages)
                metadata["page_count"] = len(reader.pages)
                style = {"heading_count": sum(bool(re.match(r"^\s*(?:[一二三四五六七八九十]+、|\d+[.、])", line)) for line in text.splitlines()),
                         "table_count": 0, "list_count": 0}
            elif extension in {".txt", ".md", ".markdown"}:
                text = data.decode("utf-8-sig")
                if "\x00" in text:
                    raise DocumentError("upload_signature_mismatch", "文本文件包含无效字符。")
                style = {"heading_count": len(re.findall(r"^#{1,6}\s", text, re.M)),
                         "table_count": len(re.findall(r"^\s*\|?\s*:?-{3,}", text, re.M)),
                         "list_count": len(re.findall(r"^\s*[-*+]\s", text, re.M))}
            else:
                raise DocumentError("unsupported_upload_format", "支持 DOCX、PDF、Markdown 和 UTF-8 文本；旧版 DOC 请另存为 DOCX。")
        except DocumentError:
            raise
        except Exception:
            raise DocumentError("document_parse_failed", "文档无法解析，请确认文件完整且格式正确。") from None
        if len(text) > 2 * 1024 * 1024:
            raise DocumentError("upload_text_limit", "文档正文超过安全限制。")
        text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text).strip()
        if not text:
            raise DocumentError("document_has_no_text", "文档没有可提取正文；扫描 PDF 请先完成文字识别。")
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        style.update({"paragraph_count": len(lines), "average_paragraph_chars": round(sum(map(len, lines)) / max(len(lines), 1)),
                      "source": "deterministic_document_analysis"})
        chunks = [{"index": index, "text": text[offset:offset + 1500]} for index, offset in enumerate(range(0, len(text), 1200))]
        metadata["title"] = lines[0].lstrip("# ")[:200]
        for key, words in (("school_stage", ["小学", "初中", "高中"]),
                           ("subject", ["语文", "数学", "英语", "物理", "化学", "生物", "历史", "地理", "道德与法治"])):
            hits = [word for word in words if word in text[:500]]
            metadata[key] = hits[0] if len(hits) == 1 else None
        from app.core.catalog import GRADES, LESSON_TYPES, STAGES, SUBJECTS
        for field, choices in [("school_stage", STAGES), ("subject", SUBJECTS)]:
            metadata[field] = next((c["value"] for c in choices if metadata.get(field) in {c["value"], c["label"]}), None)
        grade_aliases = {"grade_7": ["初一"], "grade_8": ["初二"], "grade_9": ["初三"],
                         "grade_10": ["高中一年级"], "grade_11": ["高中二年级"], "grade_12": ["高中三年级"]}
        grade_hits = [g for g in GRADES
                      if (not metadata["school_stage"] or g["school_stage"] == metadata["school_stage"])
                      and any(name in text[:500] for name in [g["label"], *grade_aliases.get(g["value"], [])])]
        if len(grade_hits) == 1 and (not metadata["school_stage"] or grade_hits[0]["school_stage"] == metadata["school_stage"]):
            metadata["grade"] = grade_hits[0]["value"]
            metadata["school_stage"] = metadata["school_stage"] or grade_hits[0]["school_stage"]
        else:
            metadata["grade"] = None
        types = [entry["value"] for entry in LESSON_TYPES if entry["label"] in text[:500]]
        metadata["lesson_type"] = types[0] if len(types) == 1 else "new_lesson"
        metadata["classification_source"] = "heuristic_requires_teacher_verification"
        return {"text": text, "metadata": metadata, "style_summary": style, "chunks": chunks}

    def _validate_zip(self, data):
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 1024 or sum(info.file_size for info in entries) > self.max_expanded_bytes:
                raise DocumentError("upload_expansion_limit", "DOCX 解压后的大小或文件数超过安全限制。")
            for info in entries:
                if info.flag_bits & 1:
                    raise DocumentError("upload_encrypted", "不支持加密 DOCX。")
                parts = Path(info.filename.replace("\\", "/")).parts
                if ".." in parts or info.filename.startswith(("/", "\\")) or ":" in info.filename:
                    raise DocumentError("upload_archive_path", "DOCX 压缩包包含不安全路径。")
            names = set(archive.namelist())
            if not {"[Content_Types].xml", "word/document.xml"} <= names:
                raise DocumentError("upload_invalid_docx", "文件不是有效的 DOCX 文档。")


def parse_upload(filename: str, data: bytes, **kwargs) -> dict:
    return DocumentAdapter(**kwargs).parse_upload(filename, data)
